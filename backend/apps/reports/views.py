"""
报表视图
"""
import logging
import os
import shutil
from collections import defaultdict
from datetime import datetime, timedelta
from decimal import Decimal
from django.utils import timezone
from django.db.models import Sum, Count, Q, F
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from apps.core.response import success_response, error_response
from apps.warehouse.models import Goods, StockIn, StockOut, Warning, StockInCorrection
from apps.warehouse import corrections as correction_service
from .models import DailyReport

logger = logging.getLogger('apps')


def _effective_quantities(stock_in_ids, moment=None):
    """按更正链重建指定入库记录在某时点的数量，返回 {id: Decimal}。"""
    if not stock_in_ids:
        return {}
    records = StockIn.objects.filter(id__in=stock_in_ids).values('id', 'initial_quantity', 'initial_batch_no')
    entries = StockInCorrection.objects.filter(
        stock_in_id__in=stock_in_ids, status='effective'
    ).order_by('seq')
    grouped = defaultdict(list)
    for entry in entries:
        grouped[entry.stock_in_id].append(entry)
    result = {}
    for record in records:
        values = correction_service.apply_entries(
            record['initial_quantity'], record['initial_batch_no'],
            grouped.get(record['id'], []), moment,
        )
        result[record['id']] = values['quantity']
    return result


def _stock_in_totals(day, moment=None):
    """
    某日入库笔数与数量合计。

    moment 为 None 时直接使用主记录当前值（默认最新状态）；
    否则只统计该时点前已登记的记录，并应用该时点前已生效的更正。
    """
    qs = StockIn.objects.filter(stock_in_time__date=day)
    if moment is not None:
        qs = qs.filter(stock_in_time__lte=moment)
    ids = list(qs.values_list('id', flat=True))
    if not ids:
        return 0, Decimal('0')
    if moment is None:
        agg = qs.aggregate(count=Count('id'), total=Sum('quantity'))
        return agg['count'] or 0, agg['total'] or Decimal('0')
    quantities = _effective_quantities(ids, moment)
    return len(ids), sum(quantities.values())


class DashboardView(APIView):
    """仪表盘数据"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        today = datetime.now().date()
        moment = correction_service.parse_moment(request.query_params.get('as_of'))

        # 货物统计
        goods_count = Goods.objects.filter(is_active=True).count()
        warning_goods_count = Goods.objects.filter(
            is_active=True,
            quantity__lte=F('warning_threshold')
        ).count()

        # 今日入库统计（as_of 时按更正链重建该时点数量）
        today_in_count, today_in_total = _stock_in_totals(today, moment)
        
        # 今日出库统计
        out_qs = StockOut.objects.filter(
            stock_out_time__date=today, status='completed'
        )
        if moment is not None:
            out_qs = out_qs.filter(stock_out_time__lte=moment)
        today_out = out_qs.aggregate(count=Count('id'), total=Sum('quantity'))

        # 待审批数量（时间点视图只统计当时已存在的待审批单）
        pending_qs = StockOut.objects.filter(status='pending')
        if moment is not None:
            pending_qs = pending_qs.filter(created_at__lte=moment)
        pending_approval_count = pending_qs.count()

        # 未读预警数量
        warning_qs = Warning.objects.filter(is_read=False)
        if moment is not None:
            warning_qs = warning_qs.filter(created_at__lte=moment)
        unread_warning_count = warning_qs.count()

        # 最近7天入库趋势
        in_trend = []
        out_trend = []
        for i in range(6, -1, -1):
            date = today - timedelta(days=i)
            in_count, in_total = _stock_in_totals(date, moment)
            out_day_qs = StockOut.objects.filter(
                stock_out_time__date=date, status='completed'
            )
            if moment is not None:
                out_day_qs = out_day_qs.filter(stock_out_time__lte=moment)
            out_data = out_day_qs.aggregate(total=Sum('quantity'))

            in_trend.append({
                'date': date.strftime('%m-%d'),
                'value': float(in_total or 0)
            })
            out_trend.append({
                'date': date.strftime('%m-%d'),
                'value': float(out_data['total'] or 0)
            })

        # 库存预警货物列表（时间点视图下库存本身无法精确回溯，仍展示当前结余）
        warning_goods = Goods.objects.filter(
            is_active=True,
            quantity__lte=F('warning_threshold')
        ).values('id', 'name', 'code', 'quantity', 'warning_threshold')[:5]

        return success_response(data={
            'goods_count': goods_count,
            'warning_goods_count': warning_goods_count,
            'today_in_count': today_in_count,
            'today_in_total': float(today_in_total or 0),
            'today_out_count': today_out['count'] or 0,
            'today_out_total': float(today_out['total'] or 0),
            'pending_approval_count': pending_approval_count,
            'unread_warning_count': unread_warning_count,
            'in_trend': in_trend,
            'out_trend': out_trend,
            'warning_goods': list(warning_goods),
            'as_of': request.query_params.get('as_of') if moment else None,
        })


class DailyReportView(APIView):
    """每日报表"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        start_date = request.query_params.get('start_date')
        end_date = request.query_params.get('end_date')
        moment = correction_service.parse_moment(request.query_params.get('as_of'))

        if not start_date:
            start_date = (datetime.now() - timedelta(days=30)).date()
        else:
            start_date = datetime.strptime(start_date, '%Y-%m-%d').date()

        if not end_date:
            end_date = datetime.now().date()
        else:
            end_date = datetime.strptime(end_date, '%Y-%m-%d').date()

        reports = []
        current_date = start_date

        while current_date <= end_date:
            # 入库统计（as_of 时按更正链重建该时点视图）
            in_count, in_total = _stock_in_totals(current_date, moment)

            # 出库统计
            out_qs = StockOut.objects.filter(
                stock_out_time__date=current_date, status='completed'
            )
            if moment is not None:
                out_qs = out_qs.filter(stock_out_time__lte=moment)
            out_data = out_qs.aggregate(
                count=Count('id'),
                total=Sum('quantity')
            )

            # 预警统计
            warning_qs = Warning.objects.filter(created_at__date=current_date)
            if moment is not None:
                warning_qs = warning_qs.filter(created_at__lte=moment)
            warning_count = warning_qs.count()

            reports.append({
                'date': current_date.strftime('%Y-%m-%d'),
                'in_count': in_count,
                'in_total': float(in_total or 0),
                'out_count': out_data['count'] or 0,
                'out_total': float(out_data['total'] or 0),
                'warning_count': warning_count
            })

            current_date += timedelta(days=1)

        return success_response(data=reports)


class ExportView(APIView):
    """数据导出"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        export_type = request.query_params.get('type', 'goods')
        moment = correction_service.parse_moment(request.query_params.get('as_of'))
        
        wb = Workbook()
        ws = wb.active
        
        # 样式定义
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='0066FF', end_color='0066FF', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        if export_type == 'goods':
            ws.title = '货物清单'
            headers = ['编码', '名称', '品类', '品种', '规格', '库存', '单位', '存放位置', '预警阈值']
            ws.append(headers)
            
            # 设置表头样式
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border
            
            # 数据
            goods_list = Goods.objects.select_related(
                'variety', 'variety__category', 'variety__category__unit'
            ).filter(is_active=True)
            for goods in goods_list:
                unit_name = ''
                if goods.variety and goods.variety.category:
                    unit_name = goods.variety.category.unit.name
                ws.append([
                    goods.code,
                    goods.name,
                    goods.variety.category.name if goods.variety else '',
                    goods.variety.name if goods.variety else '',
                    goods.specification,
                    float(goods.quantity),
                    unit_name,
                    goods.location,
                    float(goods.warning_threshold)
                ])
            
            filename = f'货物清单_{datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx'
            
        elif export_type == 'stock_in':
            ws.title = '入库记录'
            headers = ['货物编码', '货物名称', '入库数量', '单位', '批次号', '供应商', '操作人', '入库时间']
            ws.append(headers)
            
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border
            
            stock_ins_qs = StockIn.objects.select_related(
                'goods', 'goods__variety__category__unit', 'operator'
            ).all()
            if moment is not None:
                stock_ins_qs = stock_ins_qs.filter(stock_in_time__lte=moment)
            stock_ins = list(stock_ins_qs)
            entries = StockInCorrection.objects.filter(
                stock_in_id__in=[r.id for r in stock_ins], status='effective'
            ).order_by('seq')
            grouped = defaultdict(list)
            for entry in entries:
                grouped[entry.stock_in_id].append(entry)

            for record in stock_ins:
                rebuilt = correction_service.apply_entries(
                    record.initial_quantity, record.initial_batch_no,
                    grouped.get(record.id, []), moment,
                )
                unit_name = ''
                if record.goods.variety and record.goods.variety.category:
                    unit_name = record.goods.variety.category.unit.name
                ws.append([
                    record.goods.code,
                    record.goods.name,
                    float(rebuilt['quantity']),
                    unit_name,
                    rebuilt['batch_no'],
                    record.supplier,
                    record.operator.username if record.operator else '',
                    record.stock_in_time.strftime('%Y-%m-%d %H:%M:%S')
                ])

            suffix = f"（截至{request.query_params.get('as_of')}）" if moment else ''
            filename = f'入库记录{suffix}_{datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx'

        elif export_type == 'stock_in_corrections':
            ws.title = '入库更正分录'
            headers = [
                '入库记录ID', '货物编码', '货物名称', '序号', '字段', '分录类型',
                '原值', '建议值', '状态', '理由', '拒绝理由',
                '申请人', '申请时间', '批准人', '生效时间', '撤销目标序号', '下游引用快照'
            ]
            ws.append(headers)

            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border

            correction_qs = StockInCorrection.objects.select_related(
                'stock_in', 'stock_in__goods',
                'proposed_by', 'approved_by', 'target',
            ).order_by('stock_in_id', 'seq')
            for entry in correction_qs:
                ws.append([
                    entry.stock_in_id,
                    entry.stock_in.goods.code,
                    entry.stock_in.goods.name,
                    entry.seq,
                    entry.field_label,
                    entry.get_kind_display(),
                    entry.old_value,
                    entry.new_value,
                    entry.get_status_display(),
                    entry.reason,
                    entry.reject_reason,
                    entry.proposed_by.username if entry.proposed_by else '',
                    timezone.localtime(entry.proposed_at).strftime('%Y-%m-%d %H:%M:%S') if entry.proposed_at else '',
                    entry.approved_by.username if entry.approved_by else '',
                    timezone.localtime(entry.approved_at).strftime('%Y-%m-%d %H:%M:%S') if entry.approved_at else '',
                    entry.target.seq if entry.target_id else '',
                    entry.downstream_refs,
                ])

            filename = f'入库更正分录_{datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx'
            
        elif export_type == 'stock_out':
            ws.title = '出库记录'
            headers = ['货物编码', '货物名称', '出库数量', '单位', '领用人', '领用部门', '状态', '操作人', '出库时间']
            ws.append(headers)
            
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border
            
            stock_outs = StockOut.objects.select_related(
                'goods', 'goods__variety__category__unit', 'operator'
            ).all()
            for record in stock_outs:
                unit_name = ''
                if record.goods.variety and record.goods.variety.category:
                    unit_name = record.goods.variety.category.unit.name
                ws.append([
                    record.goods.code,
                    record.goods.name,
                    float(record.quantity),
                    unit_name,
                    record.receiver,
                    record.receiver_dept,
                    record.get_status_display(),
                    record.operator.username if record.operator else '',
                    record.stock_out_time.strftime('%Y-%m-%d %H:%M:%S') if record.stock_out_time else ''
                ])
            
            filename = f'出库记录_{datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx'
        else:
            return error_response(message='无效的导出类型')
        
        # 调整列宽
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2) * 1.2
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # 返回文件
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        wb.save(response)
        
        logger.info(f"Data exported: {export_type}")
        return response


class SystemMonitorView(APIView):
    """返回容器内可直接读取的轻量运行状态。"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        cpu_count = os.cpu_count() or 1
        load_1m = os.getloadavg()[0] if hasattr(os, "getloadavg") else 0.0
        disk_total, disk_used, _ = shutil.disk_usage('/')
        disk_percent = round(disk_used * 100 / disk_total, 2) if disk_total else 0
        process_count = len([name for name in os.listdir('/proc') if name.isdigit()]) if os.path.isdir('/proc') else 1
        
        return success_response(data={
            'cpu': {
                'percent': round(load_1m * 100 / cpu_count, 2),
                'count': cpu_count
            },
            'memory': {
                'total': None,
                'used': None,
                'percent': None
            },
            'disk': {
                'total': round(disk_total / (1024 ** 3), 2),
                'used': round(disk_used / (1024 ** 3), 2),
                'percent': disk_percent
            },
            'network': {
                'bytes_sent': None,
                'bytes_recv': None
            },
            'system': {
                'boot_time': None,
                'uptime': None,
                'process_count': process_count
            }
        })

from datetime import timedelta
from decimal import Decimal
from unittest import mock

from django.db import IntegrityError
from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient

from apps.authentication.backends import generate_token
from apps.authentication.models import User
from apps.reports.models import DailyReport
from .models import (
    Approval, Category, Goods, StockIn, StockInCorrection, StockOut, Unit, Variety, Warning,
)


class WarehouseFixture(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("warehouse-user", "testpass123", role="admin")
        self.client = APIClient()
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {generate_token(self.user)}")
        self.unit = Unit.objects.create(name="件", created_by=self.user)
        self.category = Category.objects.create(name="受控器材", unit=self.unit, created_by=self.user)
        self.variety = Variety.objects.create(name="记录终端", category=self.category, created_by=self.user)
        self.goods = Goods.objects.create(
            variety=self.variety,
            name="执法记录终端",
            code="DEV-001",
            quantity=Decimal("12"),
            warning_threshold=Decimal("5"),
        )


class WarehouseModelTest(WarehouseFixture):
    def test_relationship_flags(self):
        self.assertTrue(self.unit.is_linked)
        self.assertTrue(self.category.is_linked)
        self.assertTrue(self.variety.is_in_stock)
        self.assertFalse(self.goods.is_warning)

    def test_unique_unit_name(self):
        with self.assertRaises(IntegrityError):
            Unit.objects.create(name="件", created_by=self.user)

    def test_stock_records_and_approval(self):
        inbound = StockIn.objects.create(goods=self.goods, operator=self.user, quantity=Decimal("3"))
        outbound = StockOut.objects.create(
            goods=self.goods, operator=self.user, receiver="保管员", quantity=Decimal("2")
        )
        approval = Approval.objects.create(stock_out=outbound, approver=self.user)
        self.assertEqual(inbound.goods_id, self.goods.id)
        self.assertEqual(approval.status, "pending")

    def test_warning_record(self):
        warning = Warning.objects.create(goods=self.goods, type="low_stock", message="库存不足")
        self.assertFalse(warning.is_read)
        self.assertIn("执法记录终端", str(warning))


class WarehouseAPITest(WarehouseFixture):
    def test_list_units(self):
        response = self.client.get("/api/units/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["data"]["total"], 1)

    def test_create_unit_and_reject_duplicate(self):
        created = self.client.post("/api/units/", {"name": "箱"}, format="json")
        duplicate = self.client.post("/api/units/", {"name": "箱"}, format="json")
        self.assertEqual(created.status_code, 200)
        self.assertEqual(duplicate.status_code, 400)

    def test_update_linked_unit(self):
        response = self.client.put(f"/api/units/{self.unit.id}/", {"name": "台"}, format="json")
        self.assertEqual(response.status_code, 200)
        self.unit.refresh_from_db()
        self.assertEqual(self.unit.name, "台")

    def test_refuse_delete_linked_unit(self):
        response = self.client.delete(f"/api/units/{self.unit.id}/")
        self.assertEqual(response.status_code, 400)
        self.assertTrue(Unit.objects.filter(pk=self.unit.id).exists())

    def test_create_category_validates_unit(self):
        ok = self.client.post("/api/categories/", {"name": "封存介质", "unit": self.unit.id}, format="json")
        bad = self.client.post("/api/categories/", {"name": "无效分类", "unit": 99999}, format="json")
        self.assertEqual(ok.status_code, 200)
        self.assertEqual(bad.status_code, 400)

    def test_create_variety_and_duplicate_boundary(self):
        ok = self.client.post("/api/varieties/", {"name": "封存硬盘", "category": self.category.id}, format="json")
        duplicate = self.client.post("/api/varieties/", {"name": "封存硬盘", "category": self.category.id}, format="json")
        self.assertEqual(ok.status_code, 200)
        self.assertEqual(duplicate.status_code, 400)

    def test_requires_authentication(self):
        anonymous = APIClient().get("/api/units/")
        self.assertEqual(anonymous.status_code, 401)


class StockInCorrectionFixture(TestCase):
    """更正链测试基类：一名经办人 + 一名审批人 + 一条入库记录"""

    def setUp(self):
        self.operator = User.objects.create_user("operator", "testpass123", role="user")
        self.approver = User.objects.create_user("approver", "testpass123", role="user")
        self.admin = User.objects.create_user("admin-user", "testpass123", role="admin")
        self.client = APIClient()
        self.unit = Unit.objects.create(name="件", created_by=self.operator)
        self.category = Category.objects.create(name="受控器材", unit=self.unit, created_by=self.operator)
        self.variety = Variety.objects.create(name="记录终端", category=self.category, created_by=self.operator)
        self.goods = Goods.objects.create(
            variety=self.variety,
            name="执法记录终端",
            code="DEV-001",
            quantity=Decimal("100"),
            warning_threshold=Decimal("5"),
        )
        self.stock_in = StockIn.objects.create(
            goods=self.goods, operator=self.operator,
            quantity=Decimal("100"), batch_no="BATCH-01", supplier="供应商A",
        )

    def login(self, user):
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {generate_token(user)}")

    def propose(self, items, reason="录入错误", user=None):
        self.login(user or self.operator)
        return self.client.post(
            f"/api/stock-in/{self.stock_in.id}/corrections/",
            {"items": items, "reason": reason},
            format="json",
        )

    def approve(self, correction_id, user=None):
        self.login(user or self.approver)
        return self.client.post(f"/api/stock-in/corrections/{correction_id}/approve/", {}, format="json")

    def refresh(self):
        self.stock_in.refresh_from_db()
        self.goods.refresh_from_db()


class StockInCreateTest(StockInCorrectionFixture):
    def test_create_stock_in_increases_goods_quantity(self):
        self.login(self.operator)
        response = self.client.post("/api/stock-in/", {
            "goods": self.goods.id, "quantity": "25.5", "batch_no": "B-02",
        }, format="json")
        self.assertEqual(response.status_code, 200)
        self.goods.refresh_from_db()
        self.assertEqual(self.goods.quantity, Decimal("125.50"))

    def test_create_stock_in_validates_quantity(self):
        self.login(self.operator)
        for bad in ["0", "-3", "1.005", "abc"]:
            response = self.client.post("/api/stock-in/", {
                "goods": self.goods.id, "quantity": bad,
            }, format="json")
            self.assertEqual(response.status_code, 400, bad)
        self.goods.refresh_from_db()
        self.assertEqual(self.goods.quantity, Decimal("100"))

    def test_create_stock_in_rejects_unknown_goods(self):
        self.login(self.operator)
        response = self.client.post("/api/stock-in/", {
            "goods": 99999, "quantity": "1",
        }, format="json")
        self.assertEqual(response.status_code, 404)


class StockInCorrectionFlowTest(StockInCorrectionFixture):
    def test_propose_creates_pending_correction_without_touching_record(self):
        response = self.propose([
            {"field_name": "batch_no", "new_value": "BATCH-01A"},
            {"field_name": "quantity", "new_value": "80"},
        ])
        self.assertEqual(response.status_code, 200)
        data = response.json()["data"]
        self.assertEqual(data["status"], "pending")
        items = {item["field_name"]: item for item in data["items"]}
        self.assertEqual(items["batch_no"]["old_value"], "BATCH-01")
        self.assertEqual(items["batch_no"]["new_value"], "BATCH-01A")
        self.assertEqual(items["quantity"]["old_value"], "100.00")
        self.assertEqual(items["quantity"]["new_value"], "80.00")
        # 主记录在审批前保持原值
        self.refresh()
        self.assertEqual(self.stock_in.batch_no, "BATCH-01")
        self.assertEqual(self.stock_in.quantity, Decimal("100"))

    def test_approve_applies_values_and_adjusts_stock(self):
        correction_id = self.propose([
            {"field_name": "batch_no", "new_value": "BATCH-01A"},
            {"field_name": "quantity", "new_value": "80"},
        ]).json()["data"]["id"]
        response = self.approve(correction_id)
        self.assertEqual(response.status_code, 200)
        data = response.json()["data"]
        self.assertEqual(data["status"], "effective")
        self.assertEqual(data["approved_by"], self.approver.id)
        self.assertIsNotNone(data["approved_at"])
        self.refresh()
        self.assertEqual(self.stock_in.batch_no, "BATCH-01A")
        self.assertEqual(self.stock_in.quantity, Decimal("80.00"))
        # 库存按差额同步调减
        self.assertEqual(self.goods.quantity, Decimal("80.00"))

    def test_proposer_cannot_approve_own_correction(self):
        correction_id = self.propose([{"field_name": "batch_no", "new_value": "X"}]).json()["data"]["id"]
        response = self.approve(correction_id, user=self.operator)
        self.assertEqual(response.status_code, 403)
        self.refresh()
        self.assertEqual(self.stock_in.batch_no, "BATCH-01")

    def test_only_one_pending_correction_per_record(self):
        self.assertEqual(self.propose([{"field_name": "batch_no", "new_value": "X"}]).status_code, 200)
        second = self.propose([{"field_name": "batch_no", "new_value": "Y"}])
        self.assertEqual(second.status_code, 400)
        self.assertIn("待审批", second.json()["message"])

    def test_consecutive_corrections_form_ordered_chain(self):
        first_id = self.propose([{"field_name": "batch_no", "new_value": "B-2"}]).json()["data"]["id"]
        self.approve(first_id)
        second_id = self.propose([{"field_name": "batch_no", "new_value": "B-3"}]).json()["data"]["id"]
        self.approve(second_id)

        self.refresh()
        self.assertEqual(self.stock_in.batch_no, "B-3")
        chain = StockInCorrection.objects.filter(stock_in=self.stock_in).order_by("sequence")
        self.assertEqual([c.sequence for c in chain], [1, 2])
        self.assertEqual(chain[1].items.get(field_name="batch_no").old_value, "B-2")

    def test_stale_proposal_rejected_at_approval(self):
        correction_id = self.propose([{"field_name": "batch_no", "new_value": "X"}]).json()["data"]["id"]
        # 模拟记录状态在审批前被其他途径改变
        StockIn.objects.filter(pk=self.stock_in.id).update(batch_no="CHANGED")
        response = self.approve(correction_id)
        self.assertEqual(response.status_code, 409)
        self.assertEqual(StockInCorrection.objects.get(pk=correction_id).status, "pending")

    def test_reject_keeps_record_unchanged(self):
        correction_id = self.propose([{"field_name": "quantity", "new_value": "50"}]).json()["data"]["id"]
        self.login(self.approver)
        response = self.client.post(
            f"/api/stock-in/corrections/{correction_id}/reject/",
            {"remark": "依据不足"}, format="json",
        )
        self.assertEqual(response.status_code, 200)
        self.refresh()
        self.assertEqual(self.stock_in.quantity, Decimal("100"))
        correction = StockInCorrection.objects.get(pk=correction_id)
        self.assertEqual(correction.status, "rejected")
        self.assertEqual(correction.approver_remark, "依据不足")

    def test_withdraw_only_by_proposer_and_only_pending(self):
        correction_id = self.propose([{"field_name": "batch_no", "new_value": "X"}]).json()["data"]["id"]
        self.login(self.approver)
        denied = self.client.post(f"/api/stock-in/corrections/{correction_id}/withdraw/", {}, format="json")
        self.assertEqual(denied.status_code, 403)
        self.login(self.operator)
        ok = self.client.post(f"/api/stock-in/corrections/{correction_id}/withdraw/", {}, format="json")
        self.assertEqual(ok.status_code, 200)
        again = self.client.post(f"/api/stock-in/corrections/{correction_id}/withdraw/", {}, format="json")
        self.assertEqual(again.status_code, 400)
        # 撤回后可重新提交
        self.assertEqual(self.propose([{"field_name": "batch_no", "new_value": "Y"}]).status_code, 200)

    def test_processed_correction_cannot_be_approved_again(self):
        correction_id = self.propose([{"field_name": "batch_no", "new_value": "X"}]).json()["data"]["id"]
        self.approve(correction_id)
        response = self.approve(correction_id, user=self.admin)
        self.assertEqual(response.status_code, 400)

    def test_validation_errors(self):
        # 未知字段
        self.assertEqual(self.propose([{"field_name": "supplier", "new_value": "X"}]).status_code, 400)
        # 建议值与当前值一致
        self.assertEqual(self.propose([{"field_name": "batch_no", "new_value": "BATCH-01"}]).status_code, 400)
        # 数量非法
        self.assertEqual(self.propose([{"field_name": "quantity", "new_value": "0"}]).status_code, 400)
        self.assertEqual(self.propose([{"field_name": "quantity", "new_value": "1.001"}]).status_code, 400)
        # 同一字段重复提交
        self.assertEqual(self.propose([
            {"field_name": "batch_no", "new_value": "A"},
            {"field_name": "batch_no", "new_value": "B"},
        ]).status_code, 400)
        # 理由为空
        self.assertEqual(self.propose([{"field_name": "batch_no", "new_value": "A"}], reason="").status_code, 400)
        self.assertEqual(StockInCorrection.objects.count(), 0)


class StockInCorrectionReverseTest(StockInCorrectionFixture):
    def make_effective_correction(self, field="batch_no", new_value="B-2"):
        correction_id = self.propose([{"field_name": field, "new_value": new_value}]).json()["data"]["id"]
        self.approve(correction_id)
        return correction_id

    def test_reverse_restores_original_value_after_approval(self):
        correction_id = self.make_effective_correction()
        self.login(self.operator)
        response = self.client.post(
            f"/api/stock-in/corrections/{correction_id}/reverse/",
            {"reason": "更正依据有误"}, format="json",
        )
        self.assertEqual(response.status_code, 200)
        reversal = response.json()["data"]
        self.assertTrue(reversal["is_reversal"])
        self.assertEqual(reversal["reverses"], correction_id)
        item = reversal["items"][0]
        self.assertEqual(item["old_value"], "B-2")
        self.assertEqual(item["new_value"], "BATCH-01")

        self.approve(reversal["id"])
        self.refresh()
        self.assertEqual(self.stock_in.batch_no, "BATCH-01")
        target = StockInCorrection.objects.get(pk=correction_id)
        self.assertEqual(target.reversed_by.id, reversal["id"])

    def test_reverse_quantity_restores_goods_stock(self):
        correction_id = self.make_effective_correction(field="quantity", new_value="60")
        self.refresh()
        self.assertEqual(self.goods.quantity, Decimal("60"))
        self.login(self.operator)
        reversal_id = self.client.post(
            f"/api/stock-in/corrections/{correction_id}/reverse/",
            {"reason": "冲回"}, format="json",
        ).json()["data"]["id"]
        self.approve(reversal_id)
        self.refresh()
        self.assertEqual(self.stock_in.quantity, Decimal("100"))
        self.assertEqual(self.goods.quantity, Decimal("100"))

    def test_cannot_reverse_twice(self):
        correction_id = self.make_effective_correction()
        self.login(self.operator)
        self.client.post(f"/api/stock-in/corrections/{correction_id}/reverse/", {"reason": "r"}, format="json")
        # 第一笔撤销分录处于待审批时，更正单本身不算已撤销，但待审批规则会拦截
        blocked = self.client.post(f"/api/stock-in/corrections/{correction_id}/reverse/", {"reason": "r"}, format="json")
        self.assertEqual(blocked.status_code, 400)
        # 审批生效后再次撤销 -> 已被撤销
        reversal = StockInCorrection.objects.get(reverses_id=correction_id)
        self.approve(reversal.id)
        self.login(self.operator)
        again = self.client.post(f"/api/stock-in/corrections/{correction_id}/reverse/", {"reason": "r"}, format="json")
        self.assertEqual(again.status_code, 400)

    def test_reversal_entry_cannot_be_reversed(self):
        correction_id = self.make_effective_correction()
        self.login(self.operator)
        reversal_id = self.client.post(
            f"/api/stock-in/corrections/{correction_id}/reverse/", {"reason": "r"}, format="json",
        ).json()["data"]["id"]
        self.approve(reversal_id)
        self.login(self.operator)
        response = self.client.post(f"/api/stock-in/corrections/{reversal_id}/reverse/", {"reason": "r"}, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertIn("撤销分录", response.json()["message"])

    def test_cannot_reverse_when_later_correction_overwrote_field(self):
        first_id = self.make_effective_correction(field="batch_no", new_value="B-2")
        self.make_effective_correction(field="batch_no", new_value="B-3")
        self.login(self.operator)
        response = self.client.post(f"/api/stock-in/corrections/{first_id}/reverse/", {"reason": "r"}, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertIn("后续更正", response.json()["message"])

    def test_cannot_reverse_pending_or_rejected(self):
        pending_id = self.propose([{"field_name": "batch_no", "new_value": "X"}]).json()["data"]["id"]
        self.login(self.operator)
        response = self.client.post(f"/api/stock-in/corrections/{pending_id}/reverse/", {"reason": "r"}, format="json")
        self.assertEqual(response.status_code, 400)


class StockInCorrectionReferenceRuleTest(StockInCorrectionFixture):
    def make_referenced(self):
        DailyReport.objects.create(report_date=self.stock_in.stock_in_time.date())

    def test_correction_on_referenced_record_is_flagged(self):
        self.make_referenced()
        data = self.propose([{"field_name": "batch_no", "new_value": "X"}]).json()["data"]
        self.assertTrue(data["is_post_reference"])

    def test_quantity_correction_on_referenced_record_requires_admin(self):
        self.make_referenced()
        correction_id = self.propose([{"field_name": "quantity", "new_value": "90"}]).json()["data"]["id"]
        # 普通审批人拒绝
        denied = self.approve(correction_id, user=self.approver)
        self.assertEqual(denied.status_code, 403)
        self.refresh()
        self.assertEqual(self.stock_in.quantity, Decimal("100"))
        # 管理员可以批准
        ok = self.approve(correction_id, user=self.admin)
        self.assertEqual(ok.status_code, 200)
        self.refresh()
        self.assertEqual(self.stock_in.quantity, Decimal("90.00"))

    def test_batch_no_correction_on_referenced_record_needs_no_admin(self):
        self.make_referenced()
        correction_id = self.propose([{"field_name": "batch_no", "new_value": "X"}]).json()["data"]["id"]
        response = self.approve(correction_id, user=self.approver)
        self.assertEqual(response.status_code, 200)

    def test_quantity_decrease_below_consumed_stock_is_rejected(self):
        # 后续出库业务消耗了 30，库存只剩 70；把入库数量改成 10 会使库存变为负数
        self.goods.quantity = Decimal("70")
        self.goods.save(update_fields=["quantity"])
        correction_id = self.propose([{"field_name": "quantity", "new_value": "10"}]).json()["data"]["id"]
        response = self.approve(correction_id)
        self.assertEqual(response.status_code, 400)
        self.assertIn("后续出库业务消耗", response.json()["message"])
        self.refresh()
        self.assertEqual(self.stock_in.quantity, Decimal("100"))
        self.assertEqual(self.goods.quantity, Decimal("70"))
        self.assertEqual(StockInCorrection.objects.get(pk=correction_id).status, "pending")


class StockInCorrectionAtomicityTest(StockInCorrectionFixture):
    def test_failed_approval_leaves_no_partial_state(self):
        correction_id = self.propose([
            {"field_name": "batch_no", "new_value": "B-2"},
            {"field_name": "quantity", "new_value": "80"},
        ]).json()["data"]["id"]
        self.login(self.approver)
        with mock.patch.object(Goods, "save", side_effect=RuntimeError("db write failed")):
            response = self.client.post(f"/api/stock-in/corrections/{correction_id}/approve/", {}, format="json")
        self.assertEqual(response.status_code, 500)
        # 主记录、库存、更正链三者状态一致地保持原样
        self.refresh()
        self.assertEqual(self.stock_in.batch_no, "BATCH-01")
        self.assertEqual(self.stock_in.quantity, Decimal("100"))
        self.assertEqual(self.goods.quantity, Decimal("100"))
        correction = StockInCorrection.objects.get(pk=correction_id)
        self.assertEqual(correction.status, "pending")
        self.assertIsNone(correction.approved_at)
        self.assertIsNone(correction.approved_by)


class StockInAsOfViewTest(StockInCorrectionFixture):
    def setUp(self):
        super().setUp()
        # 两笔连续更正：批次号 B-2、再改为 B-3，数量改为 80
        first = self.propose([{"field_name": "batch_no", "new_value": "B-2"}]).json()["data"]["id"]
        self.approve(first)
        second = self.propose([
            {"field_name": "batch_no", "new_value": "B-3"},
            {"field_name": "quantity", "new_value": "80"},
        ]).json()["data"]["id"]
        self.approve(second)
        # 固定生效时间，保证 as_of 查询确定性
        base = timezone.now() - timedelta(hours=2)
        StockInCorrection.objects.filter(pk=first).update(approved_at=base)
        StockInCorrection.objects.filter(pk=second).update(approved_at=base + timedelta(hours=1))
        self.t0 = (base - timedelta(hours=1)).isoformat()   # 两笔更正之前
        self.t1 = (base + timedelta(minutes=30)).isoformat()  # 两笔更正之间
        self.t2 = (base + timedelta(hours=3)).isoformat()   # 两笔更正之后

    def test_default_view_shows_latest_state(self):
        self.login(self.operator)
        response = self.client.get(f"/api/stock-in/{self.stock_in.id}/")
        data = response.json()["data"]
        self.assertEqual(data["batch_no"], "B-3")
        self.assertEqual(data["quantity"], "80.00")
        self.assertEqual(data["correction_count"], 2)

    def test_as_of_reconstructs_historical_view(self):
        self.login(self.operator)
        before_any = self.client.get(f"/api/stock-in/{self.stock_in.id}/", {"as_of": self.t0}).json()["data"]
        self.assertEqual(before_any["batch_no"], "BATCH-01")
        self.assertEqual(before_any["quantity"], "100.00")

        between = self.client.get(f"/api/stock-in/{self.stock_in.id}/", {"as_of": self.t1}).json()["data"]
        self.assertEqual(between["batch_no"], "B-2")
        self.assertEqual(between["quantity"], "100.00")

        after_all = self.client.get(f"/api/stock-in/{self.stock_in.id}/", {"as_of": self.t2}).json()["data"]
        self.assertEqual(after_all["batch_no"], "B-3")
        self.assertEqual(after_all["quantity"], "80.00")

    def test_list_view_supports_as_of(self):
        self.login(self.operator)
        latest = self.client.get("/api/stock-in/").json()["data"]["list"][0]
        self.assertEqual(latest["batch_no"], "B-3")
        historical = self.client.get("/api/stock-in/", {"as_of": self.t0}).json()["data"]["list"][0]
        self.assertEqual(historical["batch_no"], "BATCH-01")
        self.assertEqual(historical["quantity"], "100.00")

    def test_invalid_as_of_rejected(self):
        self.login(self.operator)
        response = self.client.get("/api/stock-in/", {"as_of": "not-a-time"})
        self.assertEqual(response.status_code, 400)


class StockInCorrectionAuditTest(StockInCorrectionFixture):
    def test_chain_endpoint_returns_full_history(self):
        first = self.propose([{"field_name": "batch_no", "new_value": "B-2"}]).json()["data"]["id"]
        self.approve(first)
        self.propose([{"field_name": "batch_no", "new_value": "B-3"}])

        self.login(self.operator)
        chain = self.client.get(f"/api/stock-in/{self.stock_in.id}/corrections/").json()["data"]
        self.assertEqual(len(chain), 2)
        self.assertEqual([c["sequence"] for c in chain], [1, 2])
        self.assertEqual(chain[0]["status"], "effective")
        self.assertEqual(chain[1]["status"], "pending")
        self.assertEqual(chain[0]["proposed_by_name"], "operator")
        self.assertEqual(chain[0]["approved_by_name"], "approver")

    def test_audit_list_filter(self):
        correction_id = self.propose([{"field_name": "batch_no", "new_value": "B-2"}]).json()["data"]["id"]
        self.approve(correction_id)
        self.login(self.operator)
        effective = self.client.get("/api/stock-in/corrections/", {"status": "effective"}).json()["data"]
        self.assertEqual(effective["total"], 1)
        pending = self.client.get("/api/stock-in/corrections/", {"status": "pending"}).json()["data"]
        self.assertEqual(pending["total"], 0)

    def test_correction_export_explains_differences(self):
        correction_id = self.propose([
            {"field_name": "batch_no", "new_value": "B-2"},
            {"field_name": "quantity", "new_value": "80"},
        ]).json()["data"]["id"]
        self.approve(correction_id)

        self.login(self.operator)
        response = self.client.get("/api/export/", {"type": "stock_in_correction"})
        self.assertEqual(response.status_code, 200)
        from openpyxl import load_workbook
        import io
        ws = load_workbook(io.BytesIO(response.content)).active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0][4], "字段")
        exported = {(row[4], row[5], row[6]) for row in rows[1:]}
        self.assertIn(("批次号", "BATCH-01", "B-2"), exported)
        self.assertIn(("入库数量", "100.00", "80.00"), exported)

    def test_stock_in_export_marks_corrected_records(self):
        correction_id = self.propose([{"field_name": "batch_no", "new_value": "B-2"}]).json()["data"]["id"]
        self.approve(correction_id)
        self.login(self.operator)
        response = self.client.get("/api/export/", {"type": "stock_in"})
        self.assertEqual(response.status_code, 200)
        from openpyxl import load_workbook
        import io
        ws = load_workbook(io.BytesIO(response.content)).active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0][-1], "更正次数")
        self.assertEqual(rows[1][-1], 1)
